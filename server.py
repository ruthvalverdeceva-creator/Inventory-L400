import os, json, datetime as dt
from bottle import route, run, request, static_file, response, default_app
import openpyxl

XLSX = 'CEVA_Picking_Registros.xlsx'
JOURNAL = 'journal.jsonl'

def load_static():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    def tbl(wb, sheet_name):
        ws = wb[sheet_name]
        headers = [c.value for c in ws[1]]
        data = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if any(r):
                data.append(dict(zip(headers, r)))
        return data

    cant = []
    for r in tbl(wb, 'CANTIDAD'):
        cant.append({
            'PASILLO': str(r.get('PASILLO') or '').strip(),
            'Ubicación': str(r.get('Ubicación') or '').strip(),
            'Producto': str(r.get('Producto') or '').strip(),
            'EAN': str(r.get('EAN') or '').strip(),
            'Descripción producto': str(r.get('Descripción producto') or '').strip(),
            'CAMA': str(r.get('CONVERSIÓN CAMA / PQT') or r.get('CAMA') or 0),
            'UMA': str(r.get('UMA') or '')
        })

    ops = [str(r.get('NOMBRE')).strip() for r in tbl(wb, 'OPERARIOS') if r.get('NOMBRE')]
    return wb, cant, ops

wb0, CANT_DATA, OPERARIOS = load_static()

def get_state():
    # Cargar registros previos guardados en el journal
    registros = []
    if os.path.exists(JOURNAL):
        with open(JOURNAL, 'r', encoding='utf-8') as f:
            for line in f:
                try:
                    registros.append(json.loads(line))
                except:
                    continue
    return registros

@route('/')
def root():
    return static_file('index.html', root='.')

@route('/<filepath:path>')
def server_static(filepath):
    return static_file(filepath, root='.')

@route('/api/init')
def api_init():
    return {
        'operarios': OPERARIOS,
        'cantidad': CANT_DATA,
        'registros': get_state()
    }

@route('/api/guardar', method='POST')
def api_guardar():
    data = request.json
    pasillo = str(data.get('Pasillo', '')).strip().lower()
    ubicacion = str(data.get('Ubicacion', '')).strip().lower()

    # --- VALIDACIÓN DE PASILLO ---
    # Verificar si la ubicación pertenece al pasillo seleccionado
    valid_ubis = [
        item['Ubicación'].lower() for item in CANT_DATA 
        if item['PASILLO'].lower() == pasillo
    ]

    if valid_ubis and (ubicacion not in valid_ubis):
        response.status = 400
        return {
            "status": "error",
            "message": f"⚠️ ESTA UBICACIÓN ({data.get('Ubicacion')}) NO PERTENECE AL {data.get('Pasillo').upper()}"
        }

    # Si la ubicación es correcta, se registra
    registro = {
        'Fecha': dt.datetime.now().strftime('%d %B %Y %H:%M'),
        'Operario': data.get('Operario'),
        'Pasillo': data.get('Pasillo'),
        'Ubicacion': data.get('Ubicacion'),
        'Codigo': data.get('Codigo'),
        'Descripcion': data.get('Descripcion'),
        'UMA': data.get('UMA'),
        'CamaPqt': data.get('CamaPqt'),
        'Niveles': data.get('Niveles'),
        'Saldo': data.get('Saldo'),
        'Total': data.get('Total')
    }

    # Guardar en archivo journal para persistencia rápida
    with open(JOURNAL, 'a', encoding='utf-8') as f:
        f.write(json.json_dumps(registro) if hasattr(json, 'json_dumps') else json.dumps(registro) + '\n')

    return {"status": "success", "message": "Registro guardado con éxito"}

app = default_app()

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 8000))
    run(host='0.0.0.0', port=port)
